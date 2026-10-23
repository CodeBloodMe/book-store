"""
generate_bw_compact_report.py
Generates a compact, high-density, strictly Black & White academic DBMS Project Report.
Centered university headers, enriched technical parameters, no submission file name.
Embeds ER diagrams and website screenshots cleanly.
Converts cleanly to PDF via Microsoft Word native engine.
"""

import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="4" w:space="0" w:color="D1D5DB"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def build_bw_compact_docx(markdown_path, output_docx_path):
    doc = Document()
    
    # Monochrome Palette
    BLACK = RGBColor(0, 0, 0)
    DARK_GRAY = RGBColor(30, 30, 30)
    MID_GRAY = RGBColor(100, 100, 100)
    WHITE = RGBColor(255, 255, 255)

    # Configure compact margins (0.7 in)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.different_first_page_header_footer = True
        
        # Header (Black & White)
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BCSE302P DBMS Lab Project | Krish (24BCE0846) & Parv Bhardwaj (24BCE0870)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = MID_GRAY
        
        # Footer (Black & White)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("School of Computer Science & Engineering (SCOPE), VIT | Fall Semester 2026-27")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = MID_GRAY

    with open(markdown_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    base_dir = os.path.dirname(os.path.abspath(markdown_path))

    i = 0
    total_lines = len(lines)
    in_code_block = False
    code_lines = []
    in_center_div = False
    
    while i < total_lines:
        raw_line = lines[i].rstrip('\r\n')
        line = raw_line.strip()
        
        # Handle center div flags
        if '<div align="center">' in line or '<div align="center">' in raw_line:
            in_center_div = True
            i += 1
            continue
        if '</div>' in line:
            in_center_div = False
            i += 1
            continue

        # Code block fence
        if raw_line.startswith('```'):
            if in_code_block:
                code_text = "\n".join(code_lines)
                in_code_block = False
                code_lines = []
                
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F9FAFB")
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                tc_pr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:left w:val="single" w:sz="16" w:space="0" w:color="000000"/>'
                    f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
                    f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
                    f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
                    f'</w:tcBorders>'
                )
                tc_pr.append(borders)
                
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(8)
                c_run.font.color.rgb = BLACK
                
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(2)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(raw_line)
            i += 1
            continue

        # Markdown image syntax: ![caption](image_path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if img_match:
            caption = img_match.group(1).strip()
            rel_img_path = img_match.group(2).strip()
            
            # Resolve image path
            abs_img_path = os.path.normpath(os.path.join(base_dir, rel_img_path))
            if not os.path.exists(abs_img_path):
                # Try parent dir
                parent_img_path = os.path.normpath(os.path.join(base_dir, "..", rel_img_path))
                if os.path.exists(parent_img_path):
                    abs_img_path = parent_img_path
                    
            if os.path.exists(abs_img_path):
                ip = doc.add_paragraph()
                ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ip.paragraph_format.space_before = Pt(6)
                ip.paragraph_format.space_after = Pt(2)
                ip.paragraph_format.keep_with_next = True
                irun = ip.add_run()
                irun.add_picture(abs_img_path, width=Inches(5.8))
                
                if caption:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_before = Pt(0)
                    cap.paragraph_format.space_after = Pt(5)
                    crun = cap.add_run(caption)
                    crun.font.name = "Calibri"
                    crun.font.size = Pt(8.5)
                    crun.font.italic = True
                    crun.font.bold = True
                    crun.font.color.rgb = DARK_GRAY
                print(f"Successfully embedded image: {caption} from {abs_img_path}")
            else:
                print(f"Warning: Image file not found: {abs_img_path}")
            i += 1
            continue

        # Markdown tables
        if line.startswith('|') and line.endswith('|'):
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                header_raw = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                data_rows = []
                for tline in table_lines[2:]:
                    cells = [c.strip() for c in tline.split('|')[1:-1]]
                    if len(cells) == len(header_raw):
                        data_rows.append(cells)
                    elif len(cells) > len(header_raw):
                        data_rows.append(cells[:len(header_raw)])
                    else:
                        cells.extend([''] * (len(header_raw) - len(cells)))
                        data_rows.append(cells)
                        
                tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(header_raw))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="000000", sz="6", val="single")
                
                # Header row: Pure Black background, bold White text
                hdr_cells = tbl.rows[0].cells
                for col_idx, heading in enumerate(header_raw):
                    clean_h = re.sub(r'\*\*(.*?)\*\*', r'\1', heading)
                    hdr_cells[col_idx].text = clean_h
                    set_cell_background(hdr_cells[col_idx], "000000")
                    set_cell_margins(hdr_cells[col_idx], top=45, bottom=45, left=75, right=75)
                    p = hdr_cells[col_idx].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(8.5)
                        run.font.bold = True
                        run.font.color.rgb = WHITE
                        
                # Data rows: Alternating white and light grayscale
                for row_idx, rdata in enumerate(data_rows):
                    row_cells = tbl.rows[row_idx + 1].cells
                    bg_col = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx, cell_val in enumerate(rdata):
                        clean_val = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_val)
                        clean_val = clean_val.replace('<br>', '\n')
                        row_cells[col_idx].text = clean_val
                        set_cell_background(row_cells[col_idx], bg_col)
                        set_cell_margins(row_cells[col_idx], top=24, bottom=24, left=60, right=60)
                        p = row_cells[col_idx].paragraphs[0]
                        for run in p.runs:
                            run.font.name = "Calibri"
                            run.font.size = Pt(8)
                            run.font.color.rgb = DARK_GRAY
                continue

        # Page breaks
        if line == '\\newpage':
            doc.add_page_break()
            i += 1
            continue

        # Horizontal rules
        if line in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if in_center_div:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("―" * 70)
            run.font.size = Pt(7)
            run.font.color.rgb = MID_GRAY
            i += 1
            continue

        # Check for University Centered Header Elements
        is_uni_header = in_center_div or (
            "VELLORE INSTITUTE OF TECHNOLOGY" in line or
            "School of Computer Science and Engineering" in line or
            "Lab Internal Assessment Test" in line or
            "Course: Database Systems Lab" in line or
            "CHAPTERONE" in line or
            "TABLE OF CONTENTS" in line
        )

        # Headings: Pure Black, bold, tight spacing
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            if is_uni_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(14 if not is_uni_header else 15)
            run.font.bold = True
            run.font.color.rgb = BLACK
            i += 1
            continue

        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_with_next = True
            if is_uni_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11.5 if not is_uni_header else 12)
            run.font.bold = True
            run.font.color.rgb = BLACK
            i += 1
            continue

        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_with_next = True
            if is_uni_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(10 if not is_uni_header else 10.5)
            run.font.bold = True
            run.font.color.rgb = BLACK
            i += 1
            continue

        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            if is_uni_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = BLACK
            i += 1
            continue

        # Bullet lists
        if line.startswith('- ') or line.startswith('* '):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(1)
            if in_center_div:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.color.rgb = DARK_GRAY
            i += 1
            continue

        # Numbered lists
        num_match = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if num_match:
            num = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(1)
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.color.rgb = DARK_GRAY
            i += 1
            continue

        # Blockquotes (Black & White style)
        if line.startswith('>'):
            quote_text = line[1:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "F9FAFB")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="16" w:space="0" w:color="000000"/>'
                f'<w:top w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tc_pr.append(borders)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            parts = re.split(r'(\*\*.*?\*\*)', quote_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = cp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = cp.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.italic = True
                r.font.color.rgb = BLACK
            i += 1
            continue

        # Regular text
        if line:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            if in_center_div:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.color.rgb = DARK_GRAY
        i += 1

    doc.save(output_docx_path)
    print(f"[Monochrome Report With Images] Successfully generated DOCX at: {output_docx_path}")

def convert_to_pdf(docx_path, pdf_path):
    print(f"Converting {docx_path} to PDF via docx2pdf...")
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        print(f"Successfully converted to PDF at: {pdf_path}")
        return True
    except Exception as e:
        print(f"docx2pdf failed: {e}. Trying win32com...")
        try:
            import win32com.client
            word = win32com.client.Dispatch('Word.Application')
            word.Visible = False
            in_file = os.path.abspath(docx_path)
            out_file = os.path.abspath(pdf_path)
            doc = word.Documents.Open(in_file)
            doc.SaveAs(out_file, FileFormat=17)
            doc.Close()
            word.Quit()
            print(f"Successfully converted to PDF via Word COM at: {out_file}")
            return True
        except Exception as e2:
            print(f"Word COM conversion failed: {e2}")
            return False

if __name__ == '__main__':
    md_file = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\FINAL_DBMS_PROJECT_REPORT_24BCE0846.md"
    docx_file = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\dbms_package\report\24BCE0846_L25_IA9&10_Proj_KRISH.docx"
    pdf_file = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\dbms_package\report\24BCE0846_L25_IA9&10_Proj_KRISH.pdf"
    
    build_bw_compact_docx(md_file, docx_file)
    convert_to_pdf(docx_file, pdf_file)
    
    import shutil
    root_docx = r"c:\Users\Admin\Desktop\Coading\Projects\books System\24BCE0846_L25_IA9&10_Proj_KRISH.docx"
    root_pdf = r"c:\Users\Admin\Desktop\Coading\Projects\books System\24BCE0846_L25_IA9&10_Proj_KRISH.pdf"
    shutil.copy2(docx_file, root_docx)
    if os.path.exists(pdf_file):
        shutil.copy2(pdf_file, root_pdf)
    print("All centered monochrome files with embedded images generated and mirrored successfully.")
