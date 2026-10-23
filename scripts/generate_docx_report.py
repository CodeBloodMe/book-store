"""
generate_docx_report.py
Generates the comprehensive DBMS Lab Assessment Report in DOCX and converts to PDF.
Course: BCSE302P Database Systems Lab
Faculty: Dr. V. SIVAKUMAR | Slot: L25+L26
Student: Krish (24BCE0846)
"""

import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Applies clean borders to a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def build_docx(markdown_path, output_docx_path, erd_image_path=None):
    doc = Document()
    
    # Configure page geometry (Standard Letter with 0.8 in margins)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.different_first_page_header_footer = True
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BCSE302P DBMS Lab Assessment Report | 24BCE0846 Krish")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("School of Computer Science & Engineering (SCOPE), VIT | Fall Semester 2026-27")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Read markdown content
    with open(markdown_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    SLATE = RGBColor(51, 65, 85)       # #334155
    DARK_BLUE = RGBColor(15, 23, 42)   # #0F172A
    CODE_COLOR = RGBColor(30, 41, 59)  # #1E293B

    i = 0
    total_lines = len(lines)
    in_code_block = False
    code_lines = []
    
    while i < total_lines:
        line = lines[i].rstrip('\r\n')
        
        # Check code block fences
        if line.startswith('```'):
            if in_code_block:
                # Flush code block
                code_text = "\n".join(code_lines)
                in_code_block = False
                code_lines = []
                
                # Render code box as a single-cell table with background
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                
                # Code border left
                tc_pr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:left w:val="single" w:sz="18" w:space="0" w:color="3B82F6"/>'
                    f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                    f'</w:tcBorders>'
                )
                tc_pr.append(borders)
                
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = CODE_COLOR
                
                # Add tiny spacing after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Check Table format (Markdown table: lines starting with |)
        if line.startswith('|') and line.endswith('|'):
            # Collect table lines
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                # Parse markdown table
                header_raw = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                data_rows = []
                for tline in table_lines[2:]:  # skip separator line
                    cells = [c.strip() for c in tline.split('|')[1:-1]]
                    if len(cells) == len(header_raw):
                        data_rows.append(cells)
                    elif len(cells) > len(header_raw):
                        data_rows.append(cells[:len(header_raw)])
                    else:
                        cells.extend([''] * (len(header_raw) - len(cells)))
                        data_rows.append(cells)
                        
                tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(header_raw))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="4", val="single")
                
                # Header row
                hdr_cells = tbl.rows[0].cells
                for col_idx, heading in enumerate(header_raw):
                    clean_h = re.sub(r'\*\*(.*?)\*\*', r'\1', heading)
                    hdr_cells[col_idx].text = clean_h
                    set_cell_background(hdr_cells[col_idx], "1E3A8A")
                    set_cell_margins(hdr_cells[col_idx], top=120, bottom=120, left=140, right=140)
                    p = hdr_cells[col_idx].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(9.5)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        
                # Data rows
                for row_idx, rdata in enumerate(data_rows):
                    row_cells = tbl.rows[row_idx + 1].cells
                    bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx, cell_val in enumerate(rdata):
                        clean_val = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_val)
                        clean_val = clean_val.replace('<br>', '\n')
                        row_cells[col_idx].text = clean_val
                        set_cell_background(row_cells[col_idx], bg_col)
                        set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
                        p = row_cells[col_idx].paragraphs[0]
                        for run in p.runs:
                            run.font.name = "Calibri"
                            run.font.size = Pt(8.5)
                            run.font.color.rgb = DARK_BLUE
                            
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(2)
                sp.paragraph_format.space_after = Pt(4)
            continue

        # Page breaks
        if line.strip() == '\\newpage':
            doc.add_page_break()
            i += 1
            continue

        # Horizontal rules
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("_________________________________________________________________________________")
            run.font.size = Pt(6)
            run.font.color.rgb = RGBColor(203, 213, 225)
            i += 1
            continue

        # Headers
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = NAVY
            i += 1
            continue

        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = NAVY
            i += 1
            continue

        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = SLATE
            i += 1
            continue

        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = SLATE
            i += 1
            continue

        # Bullet lists
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            bullet_text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            
            # Simple markdown bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = DARK_BLUE
            i += 1
            continue

        # Numbered lists
        num_match = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if num_match:
            num = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = DARK_BLUE
            i += 1
            continue

        # Blockquotes
        if line.strip().startswith('>'):
            quote_text = line.strip()[1:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "EFF6FF")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
                f'<w:top w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tc_pr.append(borders)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            parts = re.split(r'(\*\*.*?\*\*)', quote_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = cp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = cp.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.italic = True
                r.font.color.rgb = NAVY
            i += 1
            continue

        # Regular text
        if line.strip():
            # If line is mentioning Mermaid ERD and image exists, embed ERD image
            if "Relational Entity-Relationship Diagram (Mermaid)" in line or "## 2.1 Entity-Relationship (ER) Diagram" in line:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(4)
                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        r = p.add_run(part[2:-2])
                        r.font.bold = True
                    else:
                        r = p.add_run(part)
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK_BLUE
                
                # Check if we should insert the image here
                if erd_image_path and os.path.exists(erd_image_path):
                    ip = doc.add_paragraph()
                    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    ip.paragraph_format.space_before = Pt(6)
                    ip.paragraph_format.space_after = Pt(4)
                    irun = ip.add_run()
                    irun.add_picture(erd_image_path, width=Inches(6.2))
                    
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_before = Pt(0)
                    cap.paragraph_format.space_after = Pt(8)
                    crun = cap.add_run("Figure 2.1: ChapterOne 10-Entity Relational Schema ER Diagram")
                    crun.font.name = "Calibri"
                    crun.font.size = Pt(8.5)
                    crun.font.italic = True
                    crun.font.color.rgb = SLATE
                i += 1
                continue

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = DARK_BLUE
        i += 1

    doc.save(output_docx_path)
    print(f"Successfully generated DOCX at: {output_docx_path}")

def convert_to_pdf(docx_path, pdf_path):
    print(f"Converting {docx_path} to PDF via docx2pdf...")
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        print(f"Successfully converted to PDF at: {pdf_path}")
        return True
    except Exception as e:
        print(f"docx2pdf failed: {e}. Trying win32com...")
        try:
            import win32com.client
            word = win32com.client.Dispatch('Word.Application')
            word.Visible = False
            in_file = os.path.abspath(docx_path)
            out_file = os.path.abspath(pdf_path)
            doc = word.Documents.Open(in_file)
            doc.SaveAs(out_file, FileFormat=17) # 17 = wdFormatPDF
            doc.Close()
            word.Quit()
            print(f"Successfully converted to PDF via Word COM at: {out_file}")
            return True
        except Exception as e2:
            print(f"Word COM conversion failed: {e2}")
            return False

if __name__ == '__main__':
    md_file = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\FINAL_DBMS_PROJECT_REPORT_24BCE0846.md"
    docx_file = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\dbms_package\report\24BCE0846_L25_IA9&10_Proj_KRISH.docx"
    pdf_file = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\dbms_package\report\24BCE0846_L25_IA9&10_Proj_KRISH.pdf"
    erd_img = r"c:\Users\Admin\Desktop\Coading\Projects\books System\my-books-site\dbms_package\diagrams\ChapterOne_Current_ERD.png"
    
    build_docx(md_file, docx_file, erd_img)
    convert_to_pdf(docx_file, pdf_file)
    
    # Also copy to root directory for easy access
    import shutil
    root_docx = r"c:\Users\Admin\Desktop\Coading\Projects\books System\24BCE0846_L25_IA9&10_Proj_KRISH.docx"
    root_pdf = r"c:\Users\Admin\Desktop\Coading\Projects\books System\24BCE0846_L25_IA9&10_Proj_KRISH.pdf"
    shutil.copy2(docx_file, root_docx)
    if os.path.exists(pdf_file):
        shutil.copy2(pdf_file, root_pdf)
    print("Files mirrored to project root directory successfully.")
